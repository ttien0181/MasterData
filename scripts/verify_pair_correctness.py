"""Read-only stage-1 assessment verification on raw benchmark bundles.

Truth is opened only after candidate generation and assessment. Does not train
models, write reviewer decisions, or connect to any application database.
"""
import argparse
import csv
import hashlib
import json
import os
import sys
from collections import Counter
from pathlib import Path


def mapped_sources(path, config):
    offset = config.get('header_row', 0)
    if config.get('format', 'csv') == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = book.active.iter_rows(min_row=offset + 1, values_only=True)
            headers = next(rows)
            for values in rows:
                if any(value is not None for value in values):
                    yield dict(zip(headers, values)), book.epoch
        finally:
            book.close()
    elif config.get('format', 'csv') == 'csv':
        with path.open(encoding=config.get('encoding', 'utf-8'), newline='') as stream:
            reader = csv.reader(stream, delimiter=config.get('delimiter', ','))
            for _ in range(offset):
                next(reader)
            headers = next(reader)
            for values in reader:
                if values:
                    yield dict(zip(headers, values)) if len(values) == len(headers) else None, None
    else:
        raise ValueError('Unsupported source format: ' + config['format'])


def raw_value(field, value, config, epoch):
    from datetime import date, datetime
    if field == 'dob' and config.get('date_serial') and isinstance(value, (int, float)):
        from openpyxl.utils.datetime import WINDOWS_EPOCH, from_excel
        value = from_excel(value, epoch or WINDOWS_EPOCH)
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d') if field == 'dob' else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value) if value is not None else ''


def matcher_field(field):
    return not field.startswith(('qa_', 'truth_', 'benchmark_')) and field not in {'entity_id', 'label_status', 'record_key'}


def raw_records(bundle, cap, diagnostics=None):
    import yaml
    diagnostics = diagnostics if diagnostics is not None else {}
    diagnostics['rejected_source_rows'] = 0
    mapping = bundle / 'raw/source_mapping.yaml'
    rows = []
    if mapping.exists():
        configs = yaml.safe_load(mapping.read_text(encoding='utf-8'))['files']
        for config in configs:
            path = (bundle / config['path']).resolve()
            if not path.is_relative_to(bundle.resolve()):
                raise ValueError('Unsafe input path')
            for source, epoch in mapped_sources(path, config):
                if source is None:
                    diagnostics['rejected_source_rows'] += 1
                    continue
                identity_columns = [config['columns'].get(field)
                                    for field in ('source_record_id', 'source_system')]
                if all(column and source.get(column) == column for column in identity_columns):
                    # Some exports repeat their header inside the data section.
                    diagnostics['rejected_source_rows'] += 1
                    continue
                raw = {field: raw_value(field, source.get(column), config, epoch)
                       for field, column in config['columns'].items() if matcher_field(field)}
                if not raw.get('source_record_id', '').strip() or not raw.get('source_system', '').strip():
                    # Totals/footer lines and malformed rows have no importable
                    # source identity; report their exclusion rather than invent one.
                    diagnostics['rejected_source_rows'] += 1
                    continue
                key = f"{raw['source_system']}:{raw['source_record_id']}:{config['batch_id']}"
                rows.append((key, raw))
    else:
        for path in sorted((bundle / 'input_p119').glob('*.csv')):
            with path.open(encoding='utf-8-sig', newline='') as stream:
                for source in csv.DictReader(stream):
                    # Never project qa/truth labels into a matcher record.
                    raw = {k: v for k, v in source.items() if matcher_field(k)}
                    key = f"{raw['source_system']}:{raw['source_record_id']}"
                    rows.append((key, raw))
    if not rows:
        raise ValueError('No raw input files')
    if len({key for key, _ in rows}) != len(rows):
        raise ValueError('Duplicate raw source identity')
    available = len(rows)
    rows.sort(key=lambda item: hashlib.sha256(item[0].encode()).hexdigest())
    return rows[:cap], available


def candidate_truth_metrics(outcomes, source_keys, truth):
    """Candidate-level missed work; never feed gold back into task selection."""
    known = same = outside = 0
    by_decision, by_reason = Counter(), Counter()
    for left, right, decision, required, *reason in outcomes:
        a, b = source_keys[left], source_keys[right]
        if a not in truth or b not in truth:
            continue
        known += 1
        if truth[a] == truth[b]:
            same += 1
            if not required:
                outside += 1
                by_decision[decision] += 1
                by_reason[reason[0] if reason else 'UNSPECIFIED'] += 1
    return dict(known_candidate_pairs=known, same_entity_candidate_pairs=same,
                same_entity_outside_required=outside,
                same_entity_outside_required_by_decision=dict(by_decision),
                same_entity_outside_required_by_reason=dict(by_reason),
                interpretation='candidate-level task diagnostic; not end-to-end recall')


def verify(bundle, cap):
    from masterone.mdm.engine_adapter import engine_context
    from masterone.mdm.pipeline.blocking import generate_candidates
    from masterone.mdm.pipeline.matching import value_frequencies
    from masterone.mdm.pipeline.pair_assessment import VERSION, assess_prepared_pair, assessment_profile, assessment_record
    from masterone.mdm.pipeline.standardization import standardize
    from masterone.mdm.workflow_review_rules import review_required

    diagnostics = {}
    source, available = raw_records(bundle, cap, diagnostics)
    records, source_keys = [], {}
    for key, raw in source:
        rid = hashlib.sha256(key.encode()).hexdigest()
        nf, flags = standardize(raw, raw.get('entity_type', 'UNKNOWN'))
        record = dict(id=rid, external_id=raw['source_record_id'], source_system=raw['source_system'],
                      entity_type=raw.get('entity_type', 'UNKNOWN'), raw_fields=raw,
                      normalized_fields=nf, quality_flags=flags)
        records.append(assessment_record(record))
        source_keys[rid] = key
    by_id = {r['id']: r for r in records}
    frequencies = value_frequencies(records)
    decisions, groups, transitions = Counter(), Counter(), Counter()
    floor_failures, floor_promoted, pending, outcomes = [], 0, 0, []
    evidence_hash, candidate_hash = hashlib.sha256(), hashlib.sha256()
    with engine_context(assessment_profile()):
        candidates = generate_candidates(records)
        for (left, right), _reasons in candidates.items():
            result = assess_prepared_pair(by_id[left], by_id[right], frequencies)
            candidate_hash.update(f'{left}:{right}\n'.encode())
            evidence_hash.update(json.dumps([left, right, result.get('score'), result.get('signature')],
                                            sort_keys=True).encode())
            decisions[result['decision']] += 1
            groups[result['primary_group']] += 1
            transitions[result['rule_conclusion'] + ' -> ' + result['decision']] += 1
            required = review_required('UNDECIDED', result)
            pending += required
            outcomes.append((left, right, result['decision'], required, result.get('reason')))
            if result['rule_conclusion'] == 'INSUFFICIENT_EVIDENCE':
                floor_failures.append((left, right))
                floor_promoted += result['decision'] != 'INSUFFICIENT_EVIDENCE'
    # Assessment is finished before reading gold labels.
    import pyarrow.parquet as pq
    truth_rows = pq.read_table(bundle / 'truth/records.parquet').to_pylist()
    truth = {r['record_key']: r['entity_id'] for r in truth_rows if r.get('label_status') == 'CERTAIN'}
    if not any(key in truth for key in source_keys.values()):
        # Mentor input IDs include the temporal identity in source_record_id.
        truth = {r['source_system'] + ':' + r['source_record_id']: r['entity_id']
                 for r in truth_rows if r.get('label_status') == 'CERTAIN'}
    known, same = 0, 0
    for a, b in floor_failures:
        ka, kb = source_keys[a], source_keys[b]
        if ka in truth and kb in truth:
            known += 1
            same += truth[ka] == truth[kb]
    return {'case': bundle.name, 'assessment_version': VERSION, 'records_available': available,
            'source_parsing': diagnostics,
            'records_loaded': len(records), 'partial_sample': len(records) < available,
            'sampling': 'smallest SHA256(raw source identity); no truth used',
            'truth_used_during_compute': False, 'candidates': len(candidates),
            'candidate_hash': candidate_hash.hexdigest(), 'evidence_hash': evidence_hash.hexdigest(),
            'decisions': dict(decisions), 'groups': dict(groups), 'transitions': dict(transitions),
            'mandatory_review': pending, 'floor_failures': len(floor_failures), 'floor_promoted': floor_promoted,
            'floor_failures_with_known_truth': known, 'floor_failures_same_entity': same,
            'quality_diagnostics': candidate_truth_metrics(outcomes, source_keys, truth),
            'truth_covered_records': sum(key in truth for key in source_keys.values())}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--datasets', type=Path, required=True)
    parser.add_argument('--source-tree', type=Path, default=Path(__file__).resolve().parents[1] / 'backend/src')
    parser.add_argument('--max-records', type=int, default=12000)
    parser.add_argument('--only', nargs='*')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if args.max_records <= 0:
        parser.error('--max-records must be positive')
    sys.path.insert(0, str(args.source_tree.resolve()))
    os.environ.update(SEMANTIC_ADDRESS_ENABLED='false', MDM_SCORE_WORKERS='1',
                      LLM_ENABLED='false', LLM_SEMANTIC_ENABLED='false')
    results = []
    for bundle in sorted(args.datasets.iterdir()):
        if not bundle.is_dir() or not bundle.name.startswith('TC') or args.only and bundle.name not in args.only:
            continue
        try:
            result = verify(bundle, args.max_records)
        except Exception as error:
            result = {'case': bundle.name, 'error_type': type(error).__name__, 'error': str(error)}
        results.append(result)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(results, indent=2), encoding='utf-8')
        print(json.dumps(result), flush=True)
    return int(not results or any('error_type' in result for result in results))


if __name__ == '__main__':
    raise SystemExit(main())
